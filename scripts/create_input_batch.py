#!/usr/bin/env python3
"""
create_input_batch.py — Generates standardized input batch files from the master
StayVista OTA dataset (Current OTA Property Name .xlsx), resolving missing URLs from the net.
"""

import sys
import os
import re
import csv
import argparse
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from link_resolver import resolve_link

EXCEL_PATH = os.path.join(os.path.dirname(__file__), '..', '..', 'Current OTA Property Name .xlsx')
INPUT_XLSX_50 = os.path.join(os.path.dirname(__file__), '..', 'data', 'input_properties_batch_50.xlsx')
INPUT_CSV_50 = os.path.join(os.path.dirname(__file__), '..', 'data', 'input_properties_batch_50.csv')
INPUT_XLSX_ALL = os.path.join(os.path.dirname(__file__), '..', 'data', 'input_properties_all_4ota.xlsx')
INPUT_CSV_ALL = os.path.join(os.path.dirname(__file__), '..', 'data', 'input_properties_all_4ota.csv')

LOCATIONS = [
    'Lonavala', 'Karjat', 'Shimla', 'Kasauli', 'Goa', 'Alibaug', 'Nashik', 
    'Udaipur', 'Jaipur', 'Mussoorie', 'Ooty', 'Coorg', 'Wayanad', 'Manali', 
    'Panchgani', 'Mahabaleshwar', 'Pune', 'Igatpuri', 'Dehradun', 'Khandala',
    'Nainital', 'Bhimtal', 'Kodaikanal', 'Rishikesh', 'Chandigarh', 'Mukteshwar',
    'Palghar', 'Wada', 'Alwar', 'Dharamshala', 'Kullu', 'Kasol', 'Agra', 'Srinagar'
]

PROP_TYPES = ['Villa', 'Cottage', 'Estate', 'Homestay', 'Farmhouse', 'House', 'Bungalow', 'Retreat', 'Chalet', 'Manor']

USP_PATTERNS = [
    ('Infinity Pool', r'infinity\s*pool'),
    ('Private Pool', r'private\s*pool'),
    ('Plunge Pool', r'plunge\s*pool'),
    ('Pool', r'(?<!infinity\s)(?<!private\s)(?<!plunge\s)\bpool\b'),
    ('Mountain View', r'mountain\s*view'),
    ('Valley View', r'valley\s*view'),
    ('Lake View', r'lake\s*view'),
    ('River View', r'river(?:side|\s*view)'),
    ('Heated Jacuzzi', r'heated\s*jacuzzi'),
    ('Jacuzzi', r'(?<!heated\s)jacuzzi'),
    ('Large Lawn', r'large\s*lawn'),
    ('Lawn', r'(?<!large\s)\blawn\b'),
    ('Bonfire', r'bonfire'),
    ('BBQ', r'bbq|barbecue'),
    ('Orchard', r'orchard'),
    ('Gazebo', r'gazebo')
]

def extract_attributes(sid, data):
    titles = [info['title'] for info in data['otas'].values() if info['title']]
    urls = [info['url'] for info in data['otas'].values() if info['url']]
    combined_text = " ".join(titles + urls)
    
    bhk_match = re.search(r'(\d+)\s*(?:bhk|br|bedroom)', combined_text, re.IGNORECASE)
    bhk = int(bhk_match.group(1)) if bhk_match else 3
    
    city = 'India'
    for loc in LOCATIONS:
        if re.search(r'\b' + loc + r'\b', combined_text, re.IGNORECASE):
            city = loc
            break
            
    ptype = 'Villa'
    for pt in PROP_TYPES:
        if re.search(r'\b' + pt + r'\b', combined_text, re.IGNORECASE):
            ptype = pt
            break
            
    usps = []
    for label, pat in USP_PATTERNS:
        if re.search(pat, combined_text, re.IGNORECASE):
            if label == 'Pool' and any(p in usps for p in ['Infinity Pool', 'Private Pool', 'Plunge Pool']):
                continue
            if label == 'Lawn' and 'Large Lawn' in usps:
                continue
            if label not in usps:
                usps.append(label)
    if not usps:
        usps = ['Private Pool' if ptype == 'Villa' else 'Lawn']

    name_cand = None
    for t in titles:
        m = re.search(r'(?:StayVista(?:\'s| at| \|)?\s*)([A-Za-z0-9\s]+?)(?:\s*-\s*|\s*w\/|\s*with|\s*\||\s*@|\s*\d+\s*bhk|$)', t, re.IGNORECASE)
        if m:
            clean_n = m.group(1).strip()
            if clean_n and len(clean_n) > 3 and clean_n.lower() not in ('part of stayvista', 'comfort room', 'deluxe rooms', 'rooms'):
                name_cand = clean_n
                break
    if not name_cand and titles:
        name_cand = titles[0].split('-')[0].split('|')[0].replace("StayVista's", '').replace('StayVista at', '').replace('StayVista', '').strip()
    name_cand = re.sub(r'\b(?:Villa|Cottage|Estate|House)\b', '', name_cand or '', flags=re.IGNORECASE).strip()
    if not name_cand:
        name_cand = "Vista Retreat"

    return {
        'sv_id': sid,
        'brand_name': name_cand,
        'bhk': bhk,
        'city': city,
        'property_type': ptype,
        'usps': usps
    }

def main():
    parser = argparse.ArgumentParser(description="Generate batch input from master StayVista sheet")
    parser.add_argument('--scope', choices=['50', 'all4', 'all'], default='all4', help="Scope of properties to extract")
    args = parser.parse_args()

    print(f"Loading master workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sv_map = defaultdict(lambda: {'otas': {}})

    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        sheet_clean = name.strip()
        
        for r in rows[1:]:
            if not r or len(r) < 4: continue
            ota_id, title, sv_id, url = r[0], r[1], r[2], r[3]
            status = str(r[4]).strip() if len(r) > 4 and r[4] else ''
            
            if sv_id and str(sv_id).strip() not in ('Not Found', 'None', '', '#N/A'):
                try:
                    sid = int(float(sv_id))
                except:
                    sid = str(sv_id).strip()
                
                sv_map[sid]['otas'][sheet_clean] = {
                    'id': ota_id,
                    'title': str(title).strip() if title else '',
                    'url': str(url).strip() if url else ''
                }

    all_4 = [sid for sid, d in sv_map.items() if len(d['otas']) == 4]
    
    if args.scope == '50':
        selected_sids = all_4[:50]
        target_xlsx = INPUT_XLSX_50
        target_csv = INPUT_CSV_50
    elif args.scope == 'all4':
        selected_sids = all_4
        target_xlsx = INPUT_XLSX_ALL
        target_csv = INPUT_CSV_ALL
    else: # all
        selected_sids = list(sv_map.keys())
        target_xlsx = os.path.join(os.path.dirname(__file__), '..', 'data', 'input_properties_all.xlsx')
        target_csv = os.path.join(os.path.dirname(__file__), '..', 'data', 'input_properties_all.csv')

    print(f"Selected {len(selected_sids)} properties (scope: {args.scope})")
    
    input_rows = []
    missing_resolved = 0
    
    for sid in selected_sids:
        data = sv_map[sid]
        attrs = extract_attributes(sid, data)
        for ota_name, ota_info in data['otas'].items():
            clean_ota = 'Booking.com' if 'booking' in ota_name.lower() else ('MakeMyTrip' if 'mmt' in ota_name.lower() else ota_name)
            raw_url = ota_info['url']
            
            # Resolve missing URL from net or fallback
            resolved_url, res_source = resolve_link(sid, attrs['brand_name'], attrs['city'], clean_ota, raw_url)
            if res_source in ('net_discovered', 'fallback'):
                missing_resolved += 1
                
            input_rows.append({
                'SV_ID': sid,
                'Property_Name': attrs['brand_name'],
                'BHK': attrs['bhk'],
                'Property_Type': attrs['property_type'],
                'City': attrs['city'],
                'USPs': ", ".join(attrs['usps']),
                'OTA_Platform': clean_ota,
                'Current_Title': ota_info['title'],
                'Valid_Listing_URL': resolved_url
            })

    print(f"Total rows prepared: {len(input_rows)}, Missing URLs resolved: {missing_resolved}")

    # Write CSV
    fieldnames = ['SV_ID', 'Property_Name', 'BHK', 'Property_Type', 'City', 'USPs', 'OTA_Platform', 'Current_Title', 'Valid_Listing_URL']
    os.makedirs(os.path.dirname(target_csv), exist_ok=True)
    with open(target_csv, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(input_rows)

    # Write XLSX with styling
    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = "Properties"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    border_thin = Border(left=Side(style='thin', color='CBD5E1'),
                         right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'),
                         bottom=Side(style='thin', color='CBD5E1'))

    ws.append(fieldnames)
    for col_num in range(1, len(fieldnames) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for r_idx, row in enumerate(input_rows, 2):
        for c_idx, key in enumerate(fieldnames, 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            val = row[key]
            if key == 'Valid_Listing_URL':
                cell.value = val
                cell.hyperlink = val
                cell.font = Font(name='Calibri', size=10, color='2563EB', underline='single')
            else:
                cell.value = val
                cell.font = Font(name='Calibri', size=10)
            cell.border = border_thin
            if key in ('SV_ID', 'BHK', 'Property_Type', 'City', 'OTA_Platform'):
                cell.alignment = Alignment(horizontal='center')

    ws.row_dimensions[1].height = 24
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

    out_wb.save(target_xlsx)
    print(f"Successfully saved {target_xlsx} and {target_csv} ({len(input_rows)} rows).")

if __name__ == '__main__':
    main()
